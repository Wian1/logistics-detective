import io
from pathlib import Path

import openpyxl
import pytest

from recon import ROUTES, LoadError, load_source, reconcile

DATA = Path(__file__).resolve().parent.parent / "example-data"
IVM_FILES = {"wire": "ivm-route/WIRE_DATA.xlsx", "ahk": "ivm-route/AHK_DATA.xlsx", "lsp": "ivm-route/WAREHOUSE_DATA.xlsx"}
WIT_FILES = {"wire": "witnessing-route/WIRE_DATA.xlsx", "warehouse": "witnessing-route/WAREHOUSE_DATA.xlsx"}


def run(route_name, files, overrides=None):
    route = ROUTES[route_name]
    overrides = overrides or {}
    loaded = []
    for s in route.sources:
        content = overrides.get(s.key) or (DATA / files[s.key]).read_bytes()
        loaded.append(load_source(route, s, files[s.key], content))
    return loaded, reconcile(route, loaded)


@pytest.fixture(scope="module")
def ivm():
    return run("IVM Route", IVM_FILES)


@pytest.fixture(scope="module")
def witnessing():
    return run("Witnessing Route", WIT_FILES)


def test_ivm_header_detection(ivm):
    loaded, _ = ivm
    rows = {ls.source.key: (ls.header_row, len(ls.data)) for ls in loaded}
    assert rows == {"wire": (1, 3299), "ahk": (1, 3282), "lsp": (16, 3282)}
    assert loaded[2].data["Excel Row"].iloc[0] == 17


def test_ivm_missing_bags(ivm):
    _, res = ivm
    assert len(res.missing) == 17
    assert set(res.missing["Status"]) == {"Missing from AHK, LSP"}
    assert res.missing["Bag ID"].str.startswith("K2W4_26204_04").all()


def test_ivm_lot_mismatches(ivm):
    _, res = ivm
    lots = res.mismatches[res.mismatches["Field"] == "LOT NUMBER"]
    assert lots["Disagrees with WIRE"].str.contains("AHK").sum() == 14


def test_ivm_float_noise_is_not_a_mismatch(ivm):
    _, res = ivm
    flagged = res.mismatches[(res.mismatches["Bag ID"] == "K1W3_26225_0446")
                             & (res.mismatches["Field"] == "Receiving Warehouse Weight (kg)")]
    assert flagged.empty


def test_ivm_receive_date_ignores_time(ivm):
    # WIRE has 2026-09-24 08:00, AHK has 2026-09-24 (no time), LSP has 2026-09-25.
    _, res = ivm
    row = res.mismatches[(res.mismatches["Bag ID"] == "K1W2_26212_0229")
                         & (res.mismatches["Field"] == "RECEIVE DATE")]
    assert row["Disagrees with WIRE"].tolist() == ["LSP"]


def test_ivm_offtaker_case_insensitive(ivm):
    _, res = ivm
    assert (res.mismatches["Field"] == "OFFTAKER").sum() == 0


def test_witnessing_clean(witnessing):
    loaded, res = witnessing
    assert loaded[1].header_row == 7
    assert res.summary["Bags in all files"] == 3161
    assert res.missing.empty and res.mismatches.empty and res.quality.empty


def _edit_cell(rel_path, sheet, cell, value):
    wb = openpyxl.load_workbook(DATA / rel_path)
    ws = wb[sheet] if sheet else wb.active
    ws[cell] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_mistyped_id_reported():
    # Warehouse row 8 (column C = UNIT ID.) holds K2W4_26149_0287; introduce a typo.
    content = _edit_cell(WIT_FILES["warehouse"], None, "C8", "K2W4_26149_O287")
    _, res = run("Witnessing Route", WIT_FILES, {"warehouse": content})
    status = dict(zip(res.missing["Bag ID"], res.missing["Status"]))
    assert status["K2W4_26149_O287"].startswith("Not in WIRE")
    assert status["K2W4_26149_0287"] == "Missing from Warehouse"
    assert "K2W4_26149_O287" in set(res.quality["Bag ID"])


def test_weight_off_by_one_kg_flagged():
    # Warehouse column J = RECEIVING WAREHOUSE GROSS WEIGHT; row 8 is 1761 in both files.
    content = _edit_cell(WIT_FILES["warehouse"], None, "J8", 1762)
    _, res = run("Witnessing Route", WIT_FILES, {"warehouse": content})
    assert list(res.mismatches["Field"]) == ["ZAMBIA WAREHOUSE GROSS WEIGHT (KG)"]


def test_missing_column_error():
    content = _edit_cell(WIT_FILES["warehouse"], None, "H7", "SOMETHING ELSE")
    with pytest.raises(LoadError, match="MINE SEAL"):
        run("Witnessing Route", WIT_FILES, {"warehouse": content})


def _fill_template(route, source, rows):
    from recon.templates import build_template
    wb = openpyxl.load_workbook(io.BytesIO(build_template(route, source)))
    ws = wb["Paste Data Here"]
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_witnessing_templates_round_trip(witnessing):
    from recon.templates import template_columns
    loaded, _ = witnessing
    route = ROUTES["Witnessing Route"]
    rebuilt = []
    for ls in loaded:
        field_names = [f.name for f in route.fields]
        rows = [[rec["Bag ID"], *(rec[n] for n in field_names)] for rec in ls.data.to_dict("records")]
        assert len(template_columns(route, ls.source)) == len(rows[0])
        content = _fill_template(route, ls.source, rows)
        rebuilt.append(load_source(route, ls.source, "template.xlsx", content))
    res = reconcile(route, rebuilt)
    assert [ls.header_row for ls in rebuilt] == [1, 1]
    assert res.summary["Bags in all files"] == 3161
    assert res.missing.empty and res.mismatches.empty and res.quality.empty


def test_empty_template_loads():
    from recon.templates import build_template
    route = ROUTES["Witnessing Route"]
    for s in route.sources:
        ls = load_source(route, s, "t.xlsx", build_template(route, s))
        assert ls.data.empty
