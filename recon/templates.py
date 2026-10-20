"""Blank Excel paste templates whose headers match what the loader expects."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import WEIGHT, Route, Source

DATA_SHEET = "Paste Data Here"
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def template_columns(route: Route, source: Source) -> list[tuple[str, str]]:
    """(template header, kind) for the bag ID followed by each validated field."""
    cols = [(source.key_aliases[0], "id")]
    cols += [(f.columns[source.key][0], f.kind) for f in route.fields if source.key in f.columns]
    return cols


def build_template(route: Route, source: Source) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    cols = template_columns(route, source)
    for i, (name, kind) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        dim = ws.column_dimensions[cell.column_letter]
        dim.width = max(22, len(name) + 4)
        # Text format stops Excel from mangling IDs, seals and registrations on entry.
        dim.number_format = "0" if kind == WEIGHT else "@"
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    lines = [
        f"{route.name} - {source.label} template",
        "",
        f"1. In your {source.label} source file, copy each column listed below (data rows only, no header).",
        f"2. Paste it into the matching column on the '{DATA_SHEET}' sheet, starting at row 2.",
        "   Use Paste Special > Values (Ctrl+Alt+V, V) so the source formatting doesn't come across.",
        "3. Keep every bag on one row: paste all columns from the same source rows, in the same order.",
        "4. Do not rename, reorder or delete the header row.",
        "5. Save the file and upload it in the app.",
        "",
        "Template column  <-  copy from this source column",
    ]
    lines += [f"{name}  <-  {name}" for name, _ in cols]
    for r, line in enumerate(lines, start=1):
        info.cell(row=r, column=1, value=line)
    info["A1"].font = Font(bold=True, size=14)
    info.cell(row=10, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
